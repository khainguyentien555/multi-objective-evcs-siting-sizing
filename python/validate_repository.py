from pathlib import Path
import csv, hashlib, math
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]

def fail(msg):
    raise SystemExit("VALIDATION FAILED: "+msg)

def rows(name):
    with (ROOT/'data/published'/name).open(newline='',encoding='utf-8') as f:
        return list(csv.DictReader(f))

# Required structure
required=[
 'data/published/table_I_case_study_inputs.csv',
 'data/published/table_III_charger_allocation.csv',
 'data/published/table_IV_objective_performance.csv',
 'data/published/table_V_bipso_gr_statistics.csv',
 'data/published/table_VI_dmax_sensitivity.csv',
 'data/published/objective_weight_robustness.csv',
 'data/source/result_3_methods_new.xlsx',
 'figures/source/fig01_spatial_network.png',
 'figures/source/fig02_overall_workflow.svg',
 'checksums/REFERENCE_SHA256SUMS.txt'
]
for rel in required:
    if not (ROOT/rel).is_file(): fail('missing '+rel)

# SHA-256 manifest
for line in (ROOT/'checksums/REFERENCE_SHA256SUMS.txt').read_text(encoding='utf-8').splitlines():
    if not line.strip(): continue
    expected, rel=line.split('  ',1)
    p=ROOT/rel
    if not p.is_file(): fail('manifest target missing: '+rel)
    actual=hashlib.sha256(p.read_bytes()).hexdigest()
    if actual!=expected: fail('checksum mismatch: '+rel)

# Table III totals and station statuses
r=rows('table_III_charger_allocation.csv')
total=[x for x in r if x['station']=='Total'][0]
if [int(total[k]) for k in ('bc_11kw','bc_60kw','bc_150kw')] != [105,77,6]: fail('B&C totals')
if [int(total[k]) for k in ('bipso_gr_11kw','bipso_gr_60kw','bipso_gr_150kw')] != [91,240,2]: fail('BIPSO-GR totals')
if sum(1 for x in r if x['station']!='Total' and x['bc_status']=='Open') != 8: fail('B&C open sites')
if sum(1 for x in r if x['station']!='Total' and x['bipso_gr_status']=='Open') != 6: fail('BIPSO-GR open sites')

# Source-workbook lineage: BIPSO values are direct; B&C paper values are half-up rounded
wb=load_workbook(ROOT/'data/source/result_3_methods_new.xlsx',data_only=True,read_only=True)
if not {'CPLEX','MATLAB','Sheet1'}.issubset(wb.sheetnames): fail('source workbook sheets')
ws=wb['Sheet1']
# workbook rows 5:12 are CS8..CS1
source={}
for rr in range(5,13):
    station=ws.cell(rr,2).value
    source[station]={
      'bipso':[ws.cell(rr,c).value for c in (3,4,5)],
      'bc_raw':[ws.cell(rr,c).value for c in (6,7,8)],
    }
def half_up(x): return int(math.floor(float(x)+0.5))
for x in r:
    if x['station']=='Total': continue
    s=x['station']; z=source[s]
    pub_bipso=[int(x[k]) for k in ('bipso_gr_11kw','bipso_gr_60kw','bipso_gr_150kw')]
    if [int(v) for v in z['bipso']] != pub_bipso: fail('BIPSO source mismatch '+s)
    pub_bc=[int(x[k]) for k in ('bc_11kw','bc_60kw','bc_150kw')]
    if [half_up(v) for v in z['bc_raw']] != pub_bc: fail('B&C rounding lineage mismatch '+s)

# Table IV anchors
iv={x['metric']:x for x in rows('table_IV_objective_performance.csv')}
anchors={
 'F1_hat_served_demand_term':(0.155,0.209),
 'F2_hat_load_balance':(0.018,0.033),
 'F3_hat_land_use':(0.086,0.204),
 'weighted_Z':(-0.01765,0.00883),
 'coverage_ratio':(1.0,1.0),
 'runtime_s':(0.11,0.65),
}
for k,(a,b) in anchors.items():
    if abs(float(iv[k]['bc'])-a)>1e-10 or abs(float(iv[k]['bipso_gr'])-b)>1e-10: fail('Table IV '+k)

# Table V anchors
v={x['metric']:x['value'] for x in rows('table_V_bipso_gr_statistics.csv')}
for k,a in {'runs':30,'fitness_best':958.103,'fitness_mean':975.830,'fitness_worst':994.276,'fitness_std':9.955,'gap_mean':3.78,'mean_runtime':0.465,'stabilization_iteration':37,'bc_reference_fitness':940.291}.items():
    if abs(float(v[k])-a)>1e-10: fail('Table V '+k)

# Table VI anchors
vi=rows('table_VI_dmax_sensitivity.csv')
exp=[(0.5,8,0.3283,0.498),(0.8,7,0.4532,0.784),(1.2,6,0.4727,0.996)]
for x,e in zip(vi,exp):
    got=(float(x['dmax_km']),int(x['stations_receiving_load']),float(x['average_used_distance_km']),float(x['max_used_distance_km']))
    if any(abs(float(g)-float(h))>1e-10 for g,h in zip(got,e)): fail('Table VI')

print('PUBLISHED-RESULT VALIDATION: PASS')
print(' - archival SHA-256 manifest: PASS')
print(' - source-workbook lineage / B&C rounding: PASS')
print(' - Table III charger allocation: PASS')
print(' - Table IV benchmark metrics: PASS')
print(' - Table V 30-run statistics: PASS')
print(' - Table VI Dmax sensitivity: PASS')
